"""
First-year junior invites and lunch passes (super admin only).

    python backend/tests/test_juniors.py

NOTHING HERE TOUCHES THE REAL DATABASE OR SENDS MAIL: SUPABASE_URL points at a
closed local port before any backend import, every HTTP request and SMTP
connection raises, and the zin26 helpers are replaced by an in-memory fake.
"""

import io
import os
import re
import sys
from html import unescape
from urllib.parse import unquote

os.environ["SUPABASE_URL"] = "http://127.0.0.1:9"  # closed port: nothing listens
os.environ["SUPABASE_SERVICE_ROLE_KEY"] = "test-only-not-a-key"
os.environ["SUPABASE_ANON_KEY"] = "test-only-not-a-key"
os.environ["SMTP_USER"] = ""
os.environ["SMTP_PASS"] = ""
os.environ.setdefault("AUTH_SECRET_KEY", "test-only-secret-not-a-real-key")
os.environ.setdefault("QR_SIGNING_SECRET", "test-only-qr-secret")

import requests  # noqa: E402
import smtplib  # noqa: E402


def _blocked(*_a, **_k):
    raise RuntimeError("network access is blocked in this test suite")


requests.sessions.Session.request = _blocked
smtplib.SMTP = _blocked
smtplib.SMTP_SSL = _blocked

_BACKEND = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _BACKEND not in sys.path:
    sys.path.insert(0, _BACKEND)

from services import zin26_db as db  # noqa: E402
from services import junior_service as juniors  # noqa: E402
from services import email_service  # noqa: E402
from services.zin26_db import Zin26Error  # noqa: E402

assert db.SUPABASE_URL == "http://127.0.0.1:9", "refusing to run: the real Supabase URL leaked in"

SUPER = {"id": "00000000-0000-4000-8000-000000000001", "username": "admin", "name": "Admin", "role": "SUPER_ADMIN"}


class Fake:
    def __init__(self):
        self.tables = {"juniors": [], "participants": [], "food_attendance": []}
        self.audit = []

    @staticmethod
    def _match(row, key, expr):
        value = row.get(key)
        op, _, raw = expr.partition(".")
        if op == "in":
            # Values arrive quoted and percent-encoded: in.(%22a%40x.in%22,...)
            items = re.findall(r'"((?:[^"\\]|\\.)*)"|([^,]+)', unquote(raw)[1:-1])
            return str(value) in [re.sub(r"\\(.)", r"\1", q) if q else b for q, b in items]
        raw = unquote(raw)
        if op == "is":
            return value is None if raw == "null" else str(value).lower() == raw
        if op == "eq":
            return str(value) == raw
        if op in ("like", "ilike"):
            pattern = "^" + ".*".join(re.escape(p) for p in raw.split("*")) + "$"
            return value is not None and re.match(pattern, str(value), re.I if op == "ilike" else 0) is not None
        raise AssertionError(f"fake does not understand {key}={expr}")

    def _rows(self, table, query):
        filters, order, limit = [], None, None
        for part in (query or "").split("&"):
            if "=" not in part:
                continue
            k, v = part.split("=", 1)
            if k == "select":
                continue
            if k == "order":
                col, _, d = v.partition(".")
                order = (col, d == "desc")
            elif k == "limit":
                limit = int(v)
            else:
                filters.append((k, v))
        rows = [r for r in self.tables[table] if all(self._match(r, k, e) for k, e in filters)]
        if order:
            rows.sort(key=lambda r: str(r.get(order[0]) or ""), reverse=order[1])
        return rows[:limit] if limit else rows

    def select(self, table, query="select=*"):
        return [dict(r) for r in self._rows(table, query)]

    def select_one(self, table, query):
        rows = self.select(table, query)
        return rows[0] if rows else None

    def insert(self, table, payload):
        items = payload if isinstance(payload, list) else [payload]
        if table == "juniors":
            ids = {r["junior_id"] for r in self.tables[table]}
            mails = {r["email"] for r in self.tables[table]}
            for it in items:
                if it["junior_id"] in ids or it["email"] in mails:
                    raise Zin26Error("duplicate: 23505", status=409, code="DUPLICATE")
                ids.add(it["junior_id"])
                mails.add(it["email"])
        out = []
        for it in items:
            row = {"invite_status": "PENDING", "invite_error": None, "invited_at": None, **dict(it)}
            self.tables[table].append(row)
            out.append(dict(row))
        return out

    def update(self, table, query, payload):
        rows = self._rows(table, query)
        for r in rows:
            r.update(payload)
        return [dict(r) for r in rows]

    def delete(self, table, query):
        rows = self._rows(table, query)
        self.tables[table] = [r for r in self.tables[table] if r not in rows]
        return [dict(r) for r in rows]


def harness(fn):
    def run():
        fake = Fake()
        saved = [(db, n, getattr(db, n)) for n in ("select", "select_one", "insert", "update", "delete")]
        saved.append((juniors, "log_action", juniors.log_action))
        for n in ("select", "select_one", "insert", "update", "delete"):
            setattr(db, n, getattr(fake, n))
        juniors.log_action = lambda *a, **k: fake.audit.append((a, k))
        try:
            fn(fake)
        finally:
            for module, name, value in saved:
                setattr(module, name, value)
    run.__name__ = fn.__name__
    return run


def _csv(text):
    return text.strip().encode("utf-8")


def _xlsx(rows):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- reading the sheet ---------------------------------------------------------

def test_a_csv_with_any_heading_spelling_and_food_words_is_read():
    rows, why = juniors.read_sheet("juniors.csv", _csv("""
Full Name,E-mail,Food Preference
Arun Kumar, Arun@Gmail.com ,Veg

Priya S,priya.s@gmail.com,Non Veg
Ravi,ravi@x.in,NV
Meena,meena@x.in,v
"""))
    assert not why, why
    assert [(r["name"], r["email"], r["food_preference"], r["problem"]) for r in rows] == [
        ("Arun Kumar", "arun@gmail.com", "VEG", ""),
        ("Priya S", "priya.s@gmail.com", "NON_VEG", ""),
        ("Ravi", "ravi@x.in", "NON_VEG", ""),
        ("Meena", "meena@x.in", "VEG", ""),
    ], rows
    assert [r["row"] for r in rows] == [2, 4, 5, 6], "sheet row numbers, blank rows skipped"


def test_an_xlsx_is_read_and_bad_rows_are_named_not_dropped():
    rows, why = juniors.read_sheet("J.XLSX", _xlsx([
        ["Name", "Email", "Food"],
        ["Good One", "good@x.in", "Veg"],
        ["", "noname@x.in", "Veg"],
        ["Bad Mail", "not-an-email", "Veg"],
        ["No Food", "nofood@x.in", ""],
        ["Odd Food", "odd@x.in", "Egg"],
        ["Twice", "GOOD@x.in", "Non-veg"],
    ]))
    assert not why, why
    problems = [r["problem"] for r in rows]
    assert problems[0] == ""
    assert "Name is missing" in problems[1]
    assert "does not look right" in problems[2]
    assert "Food is missing" in problems[3]
    assert "Egg" in problems[4]
    assert "Same email as row 2" in problems[5]


def test_headings_with_extra_words_around_them_are_still_understood():
    # The shape a department sheet actually arrives in.
    rows, why = juniors.read_sheet("First year CSE.xlsx", _xlsx([
        ["S.No", "Name", "Email", "Food Preference (Veg / Non veg)"],
        [1, "Abishek M", "abishekmanivel8@gmail.com", "non veg"],
        [2, "Angelina Suji S", "nglnsanthosh@gmail.com", "Veg"],
    ]))
    assert not why, why
    assert [(r["name"], r["email"], r["food_preference"], r["problem"]) for r in rows] == [
        ("Abishek M", "abishekmanivel8@gmail.com", "NON_VEG", ""),
        ("Angelina Suji S", "nglnsanthosh@gmail.com", "VEG", ""),
    ], rows

    rows, why = juniors.read_sheet("j.csv", _csv(
        "Sl,Name of the student,Student Email ID :,Veg / Non veg\nWrong,Ravi,ravi@x.in,v"))
    assert not why, why
    assert (rows[0]["name"], rows[0]["email"], rows[0]["food_preference"]) == ("Ravi", "ravi@x.in", "VEG"), rows


def test_an_exact_heading_is_never_lost_to_a_longer_one_beside_it():
    rows, why = juniors.read_sheet("j.csv", _csv(
        "Name of college,Name,Parent email,Email,Food\nGCEE,Ravi,dad@x.in,ravi@x.in,Veg"))
    assert not why, why
    assert (rows[0]["name"], rows[0]["email"]) == ("Ravi", "ravi@x.in"), rows


def test_a_sheet_with_no_food_column_at_all_is_still_refused():
    why = juniors.read_sheet("j.csv", _csv("S.No,Name,Email\n1,Ravi,ravi@x.in"))[1]
    assert "Missing: Food" in why, why


def test_a_file_that_is_not_the_juniors_sheet_is_refused_whole():
    assert juniors.read_sheet("x.pdf", b"%PDF")[1].startswith("Upload the sheet as .xlsx or .csv")
    assert "Missing: Food" in juniors.read_sheet("x.csv", _csv("Name,Email\nA,a@x.in"))[1]
    assert juniors.read_sheet("x.csv", b"")[1] == "The file is empty."
    assert "over 1 MB" in juniors.read_sheet("x.csv", b"a" * (juniors.MAX_UPLOAD_BYTES + 1))[1]
    assert "no juniors" in juniors.read_sheet("x.csv", _csv("Name,Email,Food"))[1]


# --- preview, adding, codes ---------------------------------------------------

@harness
def test_preview_flags_people_already_invited_or_already_registered(fake):
    fake.tables["juniors"].append({"junior_id": "ZIN26-J001", "email": "old@x.in", "name": "Old"})
    fake.tables["participants"].append({"user_id": "ZIN26-0123", "email": "reg@x.in"})
    res = juniors.preview("j.csv", _csv("Name,Email,Food\nNew,new@x.in,Veg\nOld,old@x.in,Veg\nReg,reg@x.in,Veg"))
    assert res["success"] and res["ready"] == 1 and res["skipped"] == 2, res
    assert "ZIN26-J001" in res["rows"][1]["problem"]
    assert "ZIN26-0123" in res["rows"][2]["problem"] and "covers lunch" in res["rows"][2]["problem"]
    assert not fake.tables["juniors"][1:], "a preview writes nothing"


CODE = re.compile(r"^ZIN26-J\d{3}$")


@harness
def test_adding_gives_each_junior_a_fresh_random_code(fake):
    fake.tables["juniors"].append({"junior_id": "ZIN26-J005", "email": "five@x.in", "name": "Five"})
    fake.tables["participants"].append({"user_id": "ZIN26-0123", "email": "reg@x.in"})
    res = juniors.add([
        {"name": "A", "email": "A@x.in", "food_preference": "VEG"},
        {"name": "B", "email": "b@x.in", "food_preference": "NON_VEG"},
        {"name": "Again", "email": "five@x.in", "food_preference": "VEG"},      # already a junior
        {"name": "Reg", "email": "reg@x.in", "food_preference": "VEG"},         # a participant
        {"name": "", "email": "c@x.in", "food_preference": "VEG"},              # rechecked: no name
        {"name": "D", "email": "a@x.in", "food_preference": "VEG"},             # same as A
    ], SUPER)
    assert res["success"], res
    assert [(a["email"], a["food_preference"]) for a in res["added"]] == [
        ("a@x.in", "VEG"), ("b@x.in", "NON_VEG")], res["added"]
    codes = [a["junior_id"] for a in res["added"]]
    assert all(CODE.match(c) for c in codes), codes
    assert len(set(codes)) == 2 and "ZIN26-J005" not in codes
    assert res["skipped"] == 2
    assert fake.audit and fake.audit[0][0][0] == "JUNIORS_ADD"


@harness
def test_sixty_juniors_get_sixty_different_codes_that_do_not_count_up(fake):
    res = juniors.add([{"name": f"Junior {i}", "email": f"j{i}@x.in", "food_preference": "VEG"}
                       for i in range(1, 61)], SUPER)
    ids = [a["junior_id"] for a in res["added"]]
    assert len(ids) == 60 and len(set(ids)) == 60 and all(CODE.match(c) for c in ids), ids
    numbers = [int(c[len("ZIN26-J"):]) for c in ids]
    # Shuffled: not handed out in sheet order, and not simply J001 ... J060.
    assert numbers != sorted(numbers) and sorted(numbers) != list(range(1, 61)), numbers[:8]


@harness
def test_a_code_already_given_out_is_never_drawn_again(fake):
    # Every number but 417 is taken: the only code left is the one handed out.
    fake.tables["juniors"] += [{"junior_id": f"ZIN26-J{n:03d}", "email": f"old{n}@x.in", "name": "Old"}
                               for n in range(1, 1000) if n != 417]
    res = juniors.add([{"name": "New", "email": "new@x.in", "food_preference": "VEG"}], SUPER)
    assert [a["junior_id"] for a in res["added"]] == ["ZIN26-J417"], res
    full = juniors.add([{"name": "More", "email": "more@x.in", "food_preference": "VEG"}], SUPER)
    assert not full["success"] and "codes left" in full["message"], full


@harness
def test_a_code_taken_at_the_same_moment_is_drawn_again(fake):
    real_insert, calls = fake.insert, []

    def racing_insert(table, payload):
        calls.append([r["junior_id"] for r in payload])
        if len(calls) == 1:  # another super admin got there first
            raise Zin26Error("duplicate key: 23505", status=409, code="DUPLICATE")
        return real_insert(table, payload)

    saved = db.insert
    db.insert = racing_insert
    try:
        res = juniors.add([{"name": "A", "email": "a@x.in", "food_preference": "VEG"},
                           {"name": "B", "email": "b@x.in", "food_preference": "VEG"}], SUPER)
    finally:
        db.insert = saved
    assert res["success"] and len(res["added"]) == 2 and len(calls) == 2, (res, calls)
    assert all(CODE.match(c) for c in calls[1])


@harness
def test_an_email_with_a_comma_or_bracket_is_refused_not_split(fake):
    rows, why = juniors.read_sheet("j.csv", _csv(
        'Name,Email,Food\nRavi,"ravi,kumar@gmail.com",Veg\nB,(b)x@x.in,Veg\nOk,ok.name+1@mail.co.in,Veg'))
    assert not why and [r["problem"] for r in rows] == [
        "This email does not look right.", "This email does not look right.", ""], rows
    res = juniors.add([{"name": "Ravi", "email": "ravi,kumar@gmail.com", "food_preference": "VEG"}], SUPER)
    assert not res["success"] and not fake.tables["juniors"], "the server checks again"


@harness
def test_known_emails_are_found_through_the_quoted_in_list(fake):
    fake.tables["juniors"].append({"junior_id": "ZIN26-J417", "email": "a.b+c@x.in", "name": "Old"})
    res = juniors.preview("j.csv", _csv("Name,Email,Food\nOld,A.B+C@x.in,Veg"))
    assert "ZIN26-J417" in res["rows"][0]["problem"], res


def test_a_huge_sheet_is_refused_before_it_is_read_into_memory():
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["Name", "Email", "Food"])
    ws.cell(row=juniors.MAX_SHEET_LINES + 50, column=1, value="far down")
    ws.cell(row=2, column=200, value="far right")
    buf = io.BytesIO()
    wb.save(buf)
    rows, why = juniors.read_sheet("j.xlsx", buf.getvalue())
    assert not rows and "more than" in why, why
    wide = Workbook()
    wide.active.append(["Name", "Email", "Food"])
    wide.active.append(["A", "a@x.in", "Veg"])
    wide.active.cell(row=2, column=500, value="stray")
    buf = io.BytesIO()
    wide.save(buf)
    assert len(juniors._table_from_upload("j.xlsx", buf.getvalue())[1]) == juniors.MAX_SHEET_COLS


# --- sending, the list, lunch -------------------------------------------------------

@harness
def test_sending_records_sent_or_failed_per_junior(fake):
    first, second = (a["junior_id"] for a in juniors.add(
        [{"name": "Ok", "email": "ok@x.in", "food_preference": "VEG"},
         {"name": "Bounce", "email": "bounce@x.in", "food_preference": "NON_VEG"}], SUPER)["added"])
    sent = []

    def fake_send(j):
        sent.append((j["email"], j["pass_code"], j["food_preference"]))
        return {"success": j["email"] != "bounce@x.in", "error": "mailbox full"}

    saved = email_service.send_junior_invite_email
    email_service.send_junior_invite_email = fake_send
    try:
        ok = juniors.send(first.lower(), SUPER)
        bad = juniors.send(second, SUPER)
    finally:
        email_service.send_junior_invite_email = saved
    assert ok["success"] and ok["invite_status"] == "SENT", ok
    assert not bad["success"] and bad["error_code"] == "EMAIL_FAILED" and "mailbox full" in bad["message"]
    assert sent == [("ok@x.in", first, "VEG"), ("bounce@x.in", second, "NON_VEG")], "the QR is the stored code"
    rows = {r["junior_id"]: r for r in fake.tables["juniors"]}
    assert rows[first]["invite_status"] == "SENT" and rows[first]["invited_at"]
    assert rows[second]["invite_status"] == "FAILED" and "mailbox full" in rows[second]["invite_error"]
    assert juniors.send("ZIN26-0123", SUPER)["error_code"] == "NOT_FOUND", "only junior codes"
    # The pass code never goes into target_id, which the dashboard feed shows to treasurers.
    invites = [a for a, _k in fake.audit if a[0].startswith("JUNIOR_INVITE")]
    assert len(invites) == 2 and all(a[2] is None for a in invites), invites


@harness
def test_send_to_all_never_mails_anyone_twice(fake):
    jid = juniors.add([{"name": "Ok", "email": "ok@x.in", "food_preference": "VEG"}], SUPER)["added"][0]["junior_id"]
    mails = []
    saved = email_service.send_junior_invite_email
    email_service.send_junior_invite_email = lambda j: mails.append(j["email"]) or {"success": True}
    try:
        assert juniors.send(jid, SUPER)["success"]
        again = juniors.send(jid, SUPER)
        assert not again["success"] and again["error_code"] == "ALREADY_SENT", again
        assert juniors.send(jid, SUPER, resend=True)["success"], "a deliberate Resend still works"
    finally:
        email_service.send_junior_invite_email = saved
    assert mails == ["ok@x.in", "ok@x.in"]


@harness
def test_two_tabs_sending_the_same_junior_mail_them_once(fake):
    jid = juniors.add([{"name": "Ok", "email": "ok@x.in", "food_preference": "VEG"}], SUPER)["added"][0]["junior_id"]
    real_select_one, mails = db.select_one, []

    def stale_read(table, query):
        row = real_select_one(table, query)
        # The other tab claims the row between this tab's read and its claim.
        fake.update("juniors", f"junior_id=eq.{jid}", {"updated_at": "2026-09-24T09:00:00+00:00"})
        return row

    saved = email_service.send_junior_invite_email
    email_service.send_junior_invite_email = lambda j: mails.append(j["email"]) or {"success": True}
    db.select_one = stale_read
    try:
        res = juniors.send(jid, SUPER)
    finally:
        db.select_one = real_select_one
        email_service.send_junior_invite_email = saved
    assert not res["success"] and res["error_code"] == "BUSY" and not mails, res


@harness
def test_an_invite_that_went_out_counts_as_sent_even_if_saving_the_status_fails(fake):
    jid = juniors.add([{"name": "Ok", "email": "ok@x.in", "food_preference": "VEG"}], SUPER)["added"][0]["junior_id"]
    real_update, calls = db.update, []

    def flaky_update(table, query, payload):
        calls.append(payload)
        if "invite_status" in payload:
            raise Zin26Error("update juniors failed: HTTP 503")
        return real_update(table, query, payload)

    saved = email_service.send_junior_invite_email
    email_service.send_junior_invite_email = lambda j: {"success": True}
    db.update = flaky_update
    try:
        res = juniors.send(jid, SUPER)
    finally:
        db.update = real_update
        email_service.send_junior_invite_email = saved
    assert res["success"] and res["invite_status"] == "SENT", res


@harness
def test_the_list_shows_who_has_eaten_and_the_food_counts(fake):
    ate, hungry = (a["junior_id"] for a in juniors.add(
        [{"name": "Ate", "email": "ate@x.in", "food_preference": "VEG"},
         {"name": "Hungry", "email": "hungry@x.in", "food_preference": "NON_VEG"}], SUPER)["added"])
    fake.tables["food_attendance"] += [
        {"id": 1, "user_id": ate, "collected_at": "2026-09-24T13:05:00+05:30"},
        {"id": 2, "user_id": "ZIN26-0099", "collected_at": "2026-09-24T13:06:00+05:30"},   # a participant
    ]
    res = juniors.listing()
    by_id = {j["junior_id"]: j for j in res["juniors"]}
    assert by_id[ate]["lunch_at"] and not by_id[hungry]["lunch_at"]
    assert res["counts"] == {"total": 2, "sent": 0, "failed": 0, "pending": 2,
                             "veg": 1, "non_veg": 1, "lunch": 1}, res["counts"]


@harness
def test_a_junior_who_has_eaten_cannot_be_removed(fake):
    ate, typo = (a["junior_id"] for a in juniors.add(
        [{"name": "Ate", "email": "ate@x.in", "food_preference": "VEG"},
         {"name": "Typo", "email": "typo@x.in", "food_preference": "VEG"}], SUPER)["added"])
    fake.tables["food_attendance"].append({"id": 1, "user_id": ate, "collected_at": "x"})
    assert juniors.remove(ate, SUPER)["error_code"] == "LUNCH_TAKEN"
    assert juniors.remove(typo, SUPER)["success"]
    removed = [a for a, _k in fake.audit if a[0] == "JUNIOR_REMOVE"]
    assert removed and removed[0][2] is None, "no pass code as the audit target"
    assert [r["junior_id"] for r in fake.tables["juniors"]] == [ate]


# --- the email --------------------------------------------------------------------

def test_the_invite_uses_the_agreed_wording_and_the_pass_code_as_the_qr():
    html = email_service.generate_junior_invite_email_html("Arun Kumar", "Veg", "ZIN26-J007")
    # The words a reader sees, not the markup around them: restyling the card
    # must not fail this, rewording it must.
    text = " ".join(unescape(re.sub(r"<[^>]+>", " ", html)).split())
    for phrase in (
        "Dear Arun Kumar",
        "Warm wishes from your seniors — the final-year students of the Department of"
        " Computer Science & Engineering.",
        "We’re happy to invite you to Zinnia 2026 , our department symposium.",
        "Thursday, 24 September 2026",
        "9:00 AM — Inauguration at the Auditorium",
        "1:00 – 2:00 PM — Lunch (show the QR pass at the food counter)",
        "Arun Kumar 🥗 VEG", "Pass code: ZIN26-J007",
        "Please download the QR code and keep it safe on your phone",
        "looking forward to having you with us",
        "With warm regards", "Team Zinnia 2026",
    ):
        assert phrase.replace(" ,", ",") in text.replace(" ,", ","), phrase
    assert 'src="cid:junior_qr"' in html, "the QR is the inline image, not a link"
    non_veg = email_service.generate_junior_invite_email_html("Arun Kumar", "Non-veg", "ZIN26-J007")
    assert "NON-VEG" in non_veg and "NON-VEG" not in html, "the food badge follows the preference"
    assert email_service._food_label("NON_VEG") == "Non-veg" and email_service._food_label("VEG") == "Veg"
    res = email_service.send_junior_invite_email({"name": "X", "email": "x@x.in", "pass_code": "ZIN26-J001"})
    assert not res["success"] and "not configured" in res["error"], "no SMTP in tests: refused, never sent"


def test_invites_reuse_a_few_gmail_connections_instead_of_logging_in_for_each():
    import threading
    from concurrent.futures import ThreadPoolExecutor

    stats = {"logins": 0, "open": 0, "max_open": 0, "sent": 0}
    lock = threading.Lock()

    class FakeSMTP:
        def __init__(self, *a, **k):
            with lock:
                stats["open"] += 1
                stats["max_open"] = max(stats["max_open"], stats["open"])
            self.busy = False
        def ehlo(self): pass
        def starttls(self): pass
        def login(self, *a):
            with lock:
                stats["logins"] += 1
        def noop(self): return (250, b"ok")
        def send_message(self, msg, to_addrs=None):
            assert not self.busy, "one connection carried two messages at once"
            self.busy = True
            import time; time.sleep(0.02)
            self.busy = False
            with lock:
                stats["sent"] += 1
        def close(self):
            with lock:
                stats["open"] -= 1

    saved = (smtplib.SMTP, email_service.SMTP_USER, email_service.SMTP_PASS, email_service.SMTP_PORT)
    smtplib.SMTP = FakeSMTP
    email_service.SMTP_USER, email_service.SMTP_PASS, email_service.SMTP_PORT = "demo@x.in", "pw", 587
    email_service._pool.clear()
    try:
        juniors_ = [{"name": f"J{i}", "email": f"j{i}@x.in", "pass_code": f"ZIN26-J{i:03d}"} for i in range(1, 11)]
        with ThreadPoolExecutor(max_workers=3) as ex:   # the page's PARALLEL
            results = list(ex.map(email_service.send_junior_invite_email, juniors_))
        assert all(r["success"] for r in results), results
        assert stats["sent"] == 10
        assert stats["logins"] <= 3, f"logged in {stats['logins']} times for 10 invites - not reused"
        assert stats["max_open"] <= email_service._POOL_MAX, stats
        assert stats["open"] == len(email_service._pool) <= email_service._POOL_MAX, "a connection was left dangling"
    finally:
        for conn, _ in list(email_service._pool):
            conn.close()
        email_service._pool.clear()
        smtplib.SMTP, email_service.SMTP_USER, email_service.SMTP_PASS, email_service.SMTP_PORT = saved


def test_the_invite_shows_a_name_as_text_and_goes_to_one_address_only():
    html = email_service.generate_junior_invite_email_html("A&B <Ravi>", "Veg", "ZIN26-J417")
    assert "A&amp;B &lt;Ravi&gt;" in html and "<Ravi>" not in html
    for bad in ("ravi,kumar@gmail.com", "ravi@x.in, other@y.in", "Ravi <ravi@x.in>", "(ravi)kumar@gmail.com"):
        res = email_service.send_junior_invite_email({"name": "R", "email": bad, "pass_code": "ZIN26-J417"})
        assert not res["success"] and "single email address" in res["error"], (bad, res)


# --- who may use it ---------------------------------------------------------------

def test_junior_routes_are_super_admin_only():
    from flask import Flask

    from middleware.auth_middleware import generate_admin_token
    from routes.admin_routes import admin_bp
    from services.participant_auth_service import generate_participant_token

    app = Flask(__name__)
    app.register_blueprint(admin_bp)
    client = app.test_client()
    saved = juniors.listing
    juniors.listing = lambda: {"success": True, "juniors": [], "counts": {}}
    try:
        url = "/api/admin/juniors"
        assert client.get(url).status_code == 401, "no token"
        participant, _ = generate_participant_token("ZIN26-1010")
        assert client.get(url, headers={"Authorization": f"Bearer {participant}"}).status_code == 401
        for role, expected in (("TREASURER", 403), ("SPOT_DESK", 403), ("EVENT_COORDINATOR", 403),
                               ("GATE_ADMIN", 403), ("FOOD_ADMIN", 403), ("SUPER_ADMIN", 200)):
            token = generate_admin_token({"id": "x", "username": role.lower(), "role": role, "name": role})
            r = client.get(url, headers={"Authorization": f"Bearer {token}"})
            assert r.status_code == expected, f"{role}: expected {expected}, got {r.status_code}"
    finally:
        juniors.listing = saved

    rules = {rule.rule for rule in app.url_map.iter_rules()}
    for path in ("/api/admin/juniors", "/api/admin/juniors/preview",
                 "/api/admin/juniors/<junior_id>/send", "/api/admin/juniors/<junior_id>"):
        assert path in rules, path


def main() -> int:
    tests = [v for k, v in sorted(globals().items()) if k.startswith("test_") and callable(v)]
    failures = []
    for fn in tests:
        try:
            fn()
            print(f"  PASS  {fn.__name__}")
        except AssertionError as e:
            failures.append((fn.__name__, str(e)))
            print(f"  FAIL  {fn.__name__}\n          {e}")
        except Exception as e:
            failures.append((fn.__name__, f"{type(e).__name__}: {e}"))
            print(f"  ERROR {fn.__name__}\n          {type(e).__name__}: {e}")
    print(f"\n{len(tests) - len(failures)}/{len(tests)} passed")
    if failures:
        print("\nFailed:")
        for name, msg in failures:
            print(f"  - {name}: {msg}")
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
